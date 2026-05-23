"""Analiz verilerini CSV / PDF / Excel olarak dışa aktarır."""
import csv
import io
from datetime import datetime
from typing import Any

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

_TR_TABLE = str.maketrans("ışğİŞĞ", "isgISG")


def _s(text: str) -> str:
    return str(text).translate(_TR_TABLE).encode("latin-1", errors="replace").decode("latin-1")


HEADERS = ["ID", "Baslik", "Kategori", "Satici", "Tutar", "Para Birimi", "Durum", "Tarih"]
HEADERS_DISPLAY = ["ID", "Başlık", "Kategori", "Satıcı", "Tutar", "Para Birimi", "Durum", "Tarih"]
HEADER_KEYS = ["id", "title", "category", "merchant", "amount", "currency", "status", "date"]


def _fmt_date(val: Any) -> str:
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    return str(val)[:10] if val else "-"


def _row_values(row: dict) -> list:
    return [
        str(row.get("id", ""))[:8],
        str(row.get("title", "")),
        str(row.get("category", "")),
        str(row.get("merchant", "")),
        f"{row.get('amount', 0):.2f}",
        str(row.get("currency", "")),
        str(row.get("status", "")),
        _fmt_date(row.get("date")),
    ]


# CSV

def export_csv(rows: list[dict], team_id: str, period: str) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(HEADERS_DISPLAY)
    for row in rows:
        writer.writerow(_row_values(row))
    return buf.getvalue().encode("utf-8-sig")  # BOM — Excel açar


# Excel

def export_excel(rows: list[dict], team_id: str, period: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = f"Harcamalar {period}"

    # Başlık satırı
    header_fill = PatternFill("solid", fgColor="1E64C8")
    header_font = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(HEADERS_DISPLAY, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Veri satırları
    for r_idx, row in enumerate(rows, 2):
        for c_idx, val in enumerate(_row_values(row), 1):
            ws.cell(row=r_idx, column=c_idx, value=val)

    # Sütun genişlikleri
    col_widths = [12, 28, 15, 20, 12, 12, 12, 14]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# PDF

class _ExportPDF(FPDF):
    def __init__(self, period: str):
        super().__init__(orientation="L", format="A4")
        self.period = period

    def header(self):
        self.set_font("Helvetica", "B", 13)
        self.cell(0, 8, _s(f"FlowTera - Harcama Analizi ({self.period})"), ln=True, align="C")
        self.set_font("Helvetica", "", 8)
        self.cell(0, 5, f"Olusturma: {datetime.now().strftime('%Y-%m-%d %H:%M')}", ln=True, align="C")
        self.ln(3)

    def footer(self):
        self.set_y(-13)
        self.set_font("Helvetica", "I", 7)
        self.set_text_color(150, 150, 150)
        self.cell(0, 8, f"Sayfa {self.page_no()}", align="C")
        self.set_text_color(0, 0, 0)


def export_pdf(rows: list[dict], team_id: str, period: str) -> bytes:
    pdf = _ExportPDF(period)
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)

    col_widths = [20, 55, 28, 38, 22, 22, 22, 27]
    row_h = 7

    # Başlık satırı
    pdf.set_fill_color(30, 100, 200)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 8)
    for h, w in zip(HEADERS, col_widths):
        pdf.cell(w, row_h, _s(h), border=1, fill=True)
    pdf.ln()

    # Veri satırları
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Helvetica", "", 7)
    for i, row in enumerate(rows):
        fill = i % 2 == 0
        if fill:
            pdf.set_fill_color(240, 245, 255)
        for val, w in zip(_row_values(row), col_widths):
            pdf.cell(w, row_h, _s(val[:30]), border=1, fill=fill)
        pdf.ln()

    pdf.ln(4)
    total_try = sum(r.get("amount", 0) for r in rows if r.get("currency") == "TRY")
    pdf.set_font("Helvetica", "B", 9)
    pdf.cell(0, 7, _s(f"Toplam Kayit: {len(rows)}   |   TRY Toplam: {total_try:,.2f}"), ln=True)

    return bytes(pdf.output())
